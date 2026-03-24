from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROPOSAL_DIR = PROJECT_ROOT / "proposal"


@dataclass(slots=True)
class MetricsBundle:
    """제안서에 넣을 핵심 정량 지표 묶음."""

    yolo_precision: float
    yolo_recall: float
    yolo_map50: float
    yolo_map50_95: float
    temporal_context_acc: float
    temporal_boundary_f1: float
    temporal_uncertainty_mae: float
    race_avg_fps: float
    race_det_avg: float
    race_frames: int
    stopcar_avg_fps: float
    stopcar_trigger_count: int
    stopcar_threshold: float
    stopcar_gap: int
    stopcar_trigger_frames: str


def _find_template() -> Path:
    templates = sorted(PROPOSAL_DIR.glob("*.pptx"))
    if not templates:
        raise FileNotFoundError("proposal 폴더에서 PPT 템플릿을 찾지 못했습니다.")
    return templates[0]


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as file:
        reader = csv.DictReader(file)
        return list(reader)


def _load_metrics() -> MetricsBundle:
    yolo_path = PROJECT_ROOT / "experiments" / "exp_004_yolo_nano_training" / "metrics.csv"
    temporal_path = PROJECT_ROOT / "experiments" / "exp_005_temporal_gru_training" / "metrics.csv"
    race_path = PROJECT_ROOT / "experiments" / "exp_006_race_video_demo" / "metrics.csv"

    yolo_rows = _read_csv_rows(yolo_path)
    yolo_best = max(yolo_rows, key=lambda row: float(row["mAP50_95_B"]))

    temporal_rows = _read_csv_rows(temporal_path)
    temporal_best = temporal_rows[0]

    race_rows = _read_csv_rows(race_path)
    race = race_rows[0]

    stopcar_dir = _latest_stopcar_run_dir()
    stopcar_summary_path = stopcar_dir / "summary.json"
    stopcar_summary = json.loads(stopcar_summary_path.read_text(encoding="utf-8"))

    reco_candidates = sorted(
        (PROJECT_ROOT / "testing_work" / "outputs").glob("*/tuning/recommendation.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if reco_candidates:
        reco = json.loads(reco_candidates[0].read_text(encoding="utf-8"))
        recommended = reco.get("recommended", {})
        threshold = float(recommended.get("threshold", stopcar_summary["tuning_params"]["threshold"]))
        gap = int(recommended.get("min_trigger_gap", stopcar_summary["tuning_params"]["min_trigger_gap"]))
        trigger_frames = str(recommended.get("trigger_frames", ""))
    else:
        tuning_params = stopcar_summary.get("tuning_params", {})
        threshold = float(tuning_params.get("threshold", 0.5))
        gap = int(tuning_params.get("min_trigger_gap", 24))
        trigger_frames = ""

    return MetricsBundle(
        yolo_precision=float(yolo_best["precision_B"]),
        yolo_recall=float(yolo_best["recall_B"]),
        yolo_map50=float(yolo_best["mAP50_B"]),
        yolo_map50_95=float(yolo_best["mAP50_95_B"]),
        temporal_context_acc=float(temporal_best["val_context_acc"]),
        temporal_boundary_f1=float(temporal_best["val_boundary_f1"]),
        temporal_uncertainty_mae=float(temporal_best["val_uncertainty_mae"]),
        race_avg_fps=float(race["avg_fps"]),
        race_det_avg=float(race["det_avg_per_frame"]),
        race_frames=int(float(race["frames"])),
        stopcar_avg_fps=float(stopcar_summary["avg_fps"]),
        stopcar_trigger_count=int(stopcar_summary["hard_brake_trigger_count"]),
        stopcar_threshold=threshold,
        stopcar_gap=gap,
        stopcar_trigger_frames=trigger_frames,
    )


def _latest_stopcar_run_dir() -> Path:
    root = PROJECT_ROOT / "testing_work" / "outputs"
    runs = [path for path in root.glob("stopcar_context_fusion_*") if path.is_dir()]
    if not runs:
        raise FileNotFoundError("testing_work/outputs에서 stopcar run 결과를 찾지 못했습니다.")
    runs.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    return runs[0]


def _latest_race_image() -> Path | None:
    image_root = PROJECT_ROOT / "artifacts" / "screenshots"
    candidates = list(image_root.glob("race_demo_*.jpg"))
    if not candidates:
        return None
    candidates.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    return candidates[0]


def _latest_stopcar_preview() -> Path | None:
    run_dir = _latest_stopcar_run_dir()
    for name in ("dashboard_preview_tuned.jpg", "dashboard_preview_latest.jpg", "dashboard_preview.jpg"):
        candidate = run_dir / name
        if candidate.exists():
            return candidate
    candidates = sorted(run_dir.glob("dashboard_preview*.jpg"))
    return candidates[-1] if candidates else None


def _set_shape_text(slide: Any, needle: str, text: str) -> bool:
    for shape in slide.shapes:
        if hasattr(shape, "text") and needle in shape.text:
            shape.text = text
            _apply_korean_font(shape)
            return True
    return False


def _apply_korean_font(shape: Any, size_pt: float | None = None) -> None:
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "맑은 고딕"
            if size_pt is not None:
                run.font.size = Pt(size_pt)


def _add_title_box(slide: Any, title: str) -> Any:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.55))
    box.text = title
    for paragraph in box.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "맑은 고딕"
            run.font.bold = True
            run.font.size = Pt(24)
            run.font.color.rgb = RGBColor(255, 255, 255)
    return box


def _fill_slide8(slide: Any, metrics: MetricsBundle) -> None:
    _add_title_box(slide, "수행 과정 및 내용 (실행 결과 시각화)")

    race_image = _latest_race_image()
    stopcar_image = _latest_stopcar_preview()

    if race_image is not None:
        slide.shapes.add_picture(str(race_image), Inches(0.5), Inches(1.05), width=Inches(6.25))
    if stopcar_image is not None:
        slide.shapes.add_picture(str(stopcar_image), Inches(6.95), Inches(1.05), width=Inches(5.85))

    left_caption = slide.shapes.add_textbox(Inches(0.55), Inches(4.55), Inches(6.2), Inches(1.3))
    left_caption.text = (
        f"race.mp4 실시간 추론 결과\n"
        f"- 처리 프레임: {metrics.race_frames} frame\n"
        f"- 평균 FPS: {metrics.race_avg_fps:.2f}\n"
        f"- 프레임당 평균 검출 수: {metrics.race_det_avg:.3f}"
    )
    _apply_korean_font(left_caption, size_pt=14)

    right_caption = slide.shapes.add_textbox(Inches(6.95), Inches(4.55), Inches(5.9), Inches(1.3))
    right_caption.text = (
        "stopcar 급정지 맥락 + 센서 융합 UI\n"
        f"- 추천 임계값: {metrics.stopcar_threshold:.2f}\n"
        f"- 최소 트리거 간격: {metrics.stopcar_gap}\n"
        f"- 급정지 트리거 수: {metrics.stopcar_trigger_count}"
    )
    _apply_korean_font(right_caption, size_pt=14)


def _fill_slide14(slide: Any) -> None:
    _add_title_box(slide, "자체점검")
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.1), Inches(12.0), Inches(5.8))
    body.text = (
        "1) 현재 달성 수준\n"
        "- 실시간 파이프라인(Spatial→Temporal→Scoring→Curation→Export) 동작 검증 완료\n"
        "- stopcar 급정지 PoC(규칙 기반) 및 임계값 튜닝 완료\n\n"
        "2) 리스크 및 보완 계획\n"
        "- 리스크: 급정지 판단이 규칙 기반이므로 일반화 한계 존재\n"
        "- 보완: 실제 급정지 라벨 데이터 추가 및 학습 기반 이벤트 분류기로 전환\n\n"
        "3) 미완료 항목\n"
        "- Jetson Orin Nano 실기기 성능/전력 실측\n"
        "- 서버 전량 전송 방식 대비 정량 비교 실험\n"
        "- IMU 등 실제 이종센서 동기화 모듈 적용"
    )
    _apply_korean_font(body, size_pt=17)


def _fill_slide15(slide: Any) -> None:
    _add_title_box(slide, "요약 및 향후 계획")
    body = slide.shapes.add_textbox(Inches(1.1), Inches(1.4), Inches(11.2), Inches(4.8))
    body.text = (
        "요약\n"
        "- 본 과제는 엣지 기반 맥락 인식 및 데이터 선별 저장 구조를 PoC 단계에서 검증하였다.\n"
        "- 차량 주행 영상에서 실시간 추론 및 급정지 맥락/센서 융합 결과를 확보하였다.\n\n"
        "향후 계획\n"
        "- Jetson 실기기 벤치마크(FPS/지연/전력)\n"
        "- 정량 비교표(성능/지연/절감률) 고도화\n"
        "- 논문 2건 목표로 실험 설계 및 결과 축적\n\n"
        "감사합니다."
    )
    _apply_korean_font(body, size_pt=18)


def _fill_table_cell(table: Any, row: int, col: int, text: str, size_pt: float = 12) -> None:
    cell = table.cell(row, col)
    cell.text = text
    tf = cell.text_frame
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.name = "맑은 고딕"
            run.font.size = Pt(size_pt)


def build_ppt() -> Path:
    template = _find_template()
    metrics = _load_metrics()
    prs = Presentation(str(template))

    # 1페이지
    _set_shape_text(prs.slides[0], "발표자", "발표자 : (작성자)")
    _set_shape_text(
        prs.slides[0],
        "창의자율과제",
        "창의자율과제\n비전인식-맥락 연계형 온디바이스 멀티모달 데이터 파이프라인",
    )

    # 3페이지 연구배경
    _set_shape_text(
        prs.slides[2],
        "창의자율과제를 통해",
        "Physical AI 확산으로 로봇/모빌리티가 현실 맥락을 실시간 이해해야 하는 요구가 증가함.\n"
        "기존 방식(전량 서버 전송 후 추론)은 대역폭/지연/비용 부담이 크며,\n"
        "맥락 중심 학습데이터를 자동 축적하기 어려움.",
    )
    _set_shape_text(prs.slides[2], "표지, 목차 등 제외", "작성 범위: 현재 단계는 PoC(구조 검증) 결과 기준으로 작성")

    # 4페이지 연구 필요성
    _set_shape_text(
        prs.slides[3],
        "연구 및 기술개발의 필요성",
        "단일 프레임 객체 탐지만으로는 연속 상황의 맥락(Context)과 이벤트 경계를 설명하기 어려움.\n"
        "본 과제는 엣지에서 중요 이벤트만 선별 저장하여 데이터셋 구축 비용을 줄이고,\n"
        "향후 VLA 학습용 정렬 멀티모달 데이터 인프라를 구축하는 데 목적이 있음.\n\n"
        "적용 도메인: 차량 주행 영상(검증용), 향후 로봇/생체신호 연동 확장",
    )

    # 5페이지 연구 목표
    _set_shape_text(
        prs.slides[4],
        "창의자율과제를 통해",
        "목표 1) 실시간 카메라 입력 기반 맥락 인식 파이프라인 구현\n"
        "목표 2) 중요도 기반 clip/keyframe 저장 정책으로 데이터 선별\n"
        "목표 3) frame/event/aligned 스키마로 학습 재사용 가능한 데이터셋 export\n\n"
        "정량 지표(현재):\n"
        f"- YOLO best mAP50-95={metrics.yolo_map50_95:.4f}, precision={metrics.yolo_precision:.4f}, recall={metrics.yolo_recall:.4f}\n"
        f"- Temporal val_context_acc={metrics.temporal_context_acc:.4f}, val_boundary_f1={metrics.temporal_boundary_f1:.4f}",
    )

    # 6페이지 추진체계
    _set_shape_text(
        prs.slides[5],
        "창의자율과제 지도교수",
        "추진 체계(모듈):\n"
        "1) Spatial Encoder: YOLO-nano 기반 특징 벡터화\n"
        "2) Temporal Encoder: GRU 기반 맥락/경계 추정\n"
        "3) Scoring: context/boundary/uncertainty/novelty 결합\n"
        "4) Auto-curation: ring buffer 기반 clip/keyframe/discard 정책\n"
        "5) Export: frame/event/aligned JSONL 저장\n\n"
        "환경 분리:\n"
        "- 학습/실험: RTX 3080 Ti (CUDA 11.8)\n"
        "- 배포/추론 타깃: Jetson Orin Nano 8GB",
    )

    # 7페이지 추진계획
    _set_shape_text(
        prs.slides[6],
        "창의자율과제 추진 일정",
        "완료(2026-03-18 기준)\n"
        "- 프로젝트 구조/문서화 및 실험 관리 체계 구축\n"
        "- YOLO 학습/연동, Temporal GRU 학습/추론 분리 구현\n"
        "- race.mp4 실시간 검증, stopcar 급정지 맥락 PoC 및 UI 구현\n\n"
        "후속 계획\n"
        "- Jetson 실기기 성능/전력 측정\n"
        "- 서버 전량 전송 대비 정량 비교 실험\n"
        "- 실센서(IMU 등) 동기화 확장 및 논문 2건 준비",
    )

    # 8페이지 수행 과정/결과 시각화
    _fill_slide8(prs.slides[7], metrics)

    # 9페이지 연구성과(논문)
    slide9 = prs.slides[8]
    _set_shape_text(
        slide9,
        "정량성과 관련, 별도/특이사항",
        f"정량 성과 요약: YOLO mAP50-95={metrics.yolo_map50_95:.4f}, "
        f"Temporal context_acc={metrics.temporal_context_acc:.4f}, "
        f"race FPS={metrics.race_avg_fps:.2f}",
    )
    _set_shape_text(
        slide9,
        "정량성과 논문/특허/대외수상 중",
        "본 제안서는 논문 성과 중심으로 작성(특허/대외수상은 후속 단계에서 추진).",
    )
    table9 = slide9.shapes[0].table
    _fill_table_cell(table9, 1, 1, "엣지 맥락 인식 기반 멀티모달 데이터 마이닝 파이프라인")
    _fill_table_cell(table9, 2, 1, "국내/국제 AI·로보틱스 학술대회(투고 예정)")
    _fill_table_cell(table9, 3, 1, "비SCI(학술대회) / 확장저널 검토")
    _fill_table_cell(table9, 4, 1, "n.n")
    _fill_table_cell(table9, 5, 1, "100% (학생 주도)")
    _fill_table_cell(table9, 6, 1, "2026.10 (예정)")
    _fill_table_cell(table9, 7, 1, "2026.12 (예정)")

    # 10/11페이지는 특허/수상 미정으로 명시
    _set_shape_text(prs.slides[9], "정량성과 관련, 별도/특이사항", "특허 성과는 후속 단계에서 추진 (현재 단계 해당 없음)")
    table10 = prs.slides[9].shapes[0].table
    for r in range(1, len(table10.rows)):
        _fill_table_cell(table10, r, 1, "해당 없음(현재 단계)")

    _set_shape_text(prs.slides[10], "정량성과 관련, 별도/특이사항", "대외수상 성과는 후속 단계에서 추진 (현재 단계 해당 없음)")
    table11 = prs.slides[10].shapes[0].table
    for r in range(1, len(table11.rows)):
        _fill_table_cell(table11, r, 1, "해당 없음(현재 단계)")

    # 12페이지 기대효과
    _set_shape_text(
        prs.slides[11],
        "앞서 서술한 정량실적 외",
        "기대효과\n"
        "1) 엣지 기반 자동 데이터 선별로 저장/전송 비용 절감 가능성 검증\n"
        "2) 맥락 연계형 스키마(frame/event/aligned)로 VLA 학습 데이터 인프라 기반 확보\n"
        "3) 차량 영상 기준 급정지 맥락 PoC 검증\n"
        f"   - 추천 임계값: {metrics.stopcar_threshold:.2f}, 최소 간격: {metrics.stopcar_gap}, 트리거 프레임: {metrics.stopcar_trigger_frames or 'N/A'}\n"
        f"   - 재검증 결과: trigger={metrics.stopcar_trigger_count}, avg_fps={metrics.stopcar_avg_fps:.2f}\n"
        "4) 한계: Jetson 실측/전력/서버대비 비교는 후속 과제",
    )

    # 13페이지 예산
    table13 = prs.slides[12].shapes[0].table
    _fill_table_cell(table13, 1, 1, "학생인건비")
    _fill_table_cell(table13, 1, 2, "연구참여 6개월")
    _fill_table_cell(table13, 1, 3, "1,800,000")
    _fill_table_cell(table13, 2, 1, "사무용품/실험소모품")
    _fill_table_cell(table13, 2, 2, "문서화/기록용")
    _fill_table_cell(table13, 2, 3, "600,000")
    _fill_table_cell(table13, 3, 1, "분석/검증비")
    _fill_table_cell(table13, 3, 2, "평가/튜닝 반복")
    _fill_table_cell(table13, 3, 3, "1,200,000")
    _fill_table_cell(table13, 4, 1, "논문게재/학회등록")
    _fill_table_cell(table13, 4, 2, "2건 목표")
    _fill_table_cell(table13, 4, 3, "1,000,000")
    _fill_table_cell(table13, 5, 1, "데이터 처리/저장 재료비")
    _fill_table_cell(table13, 5, 2, "저장장치/백업")
    _fill_table_cell(table13, 5, 3, "1,400,000")
    _fill_table_cell(table13, 6, 3, "6,000,000")
    _set_shape_text(
        prs.slides[12],
        "예산 관련, 별도/특이사항",
        "예산 관련 특이사항: GPU/개발장비는 기존 보유 자원(RTX 3080 Ti) 활용, "
        "Jetson 실기기 검증 비용은 후속 실험 단계에서 집행",
    )

    # 14/15페이지 추가 작성
    _fill_slide14(prs.slides[13])
    _fill_slide15(prs.slides[14])

    output = PROPOSAL_DIR / f"창의자율과제_선정평가_발표자료_작성본_{datetime.now().strftime('%Y%m%d')}.pptx"
    prs.save(str(output))
    return output


def main() -> None:
    output = build_ppt()
    manifest = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "output_pptx": str(output),
    }
    manifest_path = PROPOSAL_DIR / "proposal_generation_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
