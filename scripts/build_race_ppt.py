from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def latest_race_summary(root: Path) -> tuple[dict, Path]:
    candidates: list[Path] = []
    runs_root = root / "outputs" / "runs"
    for run_dir in runs_root.glob("race_video_demo_*"):
        summary_path = run_dir / "race_demo_summary.json"
        if summary_path.exists():
            candidates.append(summary_path)
    if not candidates:
        raise FileNotFoundError("race_demo_summary.json을 찾지 못했습니다.")
    latest = sorted(candidates, key=lambda p: p.stat().st_mtime, reverse=True)[0]
    return json.loads(latest.read_text(encoding="utf-8")), latest


def load_yolo_metrics(root: Path) -> tuple[dict, dict]:
    base = json.loads(
        (root / "experiments" / "exp_004_yolo_nano_training" / "runs" / "yolov8n_e3_frac002_gpu" / "train_summary.json").read_text(encoding="utf-8")
    )
    retrain = json.loads(
        (root / "experiments" / "exp_004_yolo_nano_training" / "runs" / "yolov8n_e20_frac01_gpu" / "train_summary.json").read_text(encoding="utf-8")
    )
    return base, retrain


def load_temporal_metrics(root: Path) -> dict:
    return json.loads(
        (root / "experiments" / "exp_005_temporal_gru_training" / "runs" / "gru_e8_vehicle" / "train_summary.json").read_text(encoding="utf-8")
    )


def set_title_style(shape, size: int = 34) -> None:
    p = shape.text_frame.paragraphs[0]
    p.font.name = "Malgun Gothic"
    p.font.bold = True
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(22, 32, 56)


def set_body_style(text_frame, size: int = 20) -> None:
    for para in text_frame.paragraphs:
        para.font.name = "Malgun Gothic"
        para.font.size = Pt(size)
        para.font.color.rgb = RGBColor(36, 46, 70)


def add_title_slide(prs: Presentation, date_text: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "비전인식-맥락 연계형 파이프라인\nrace.mp4 실시간 검증 발표"
    set_title_style(slide.shapes.title, 38)

    sub = slide.placeholders[1]
    sub.text = f"프로젝트 발표 요약\n{date_text}"
    set_body_style(sub.text_frame, 20)


def add_bullet_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    set_title_style(slide.shapes.title, 30)

    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for idx, item in enumerate(bullets):
        para = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        para.text = item
        para.level = 0
    set_body_style(body, 20)


def add_metrics_slide(prs: Presentation, yolo_base: dict, yolo_retrain: dict, temporal: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "핵심 학습 결과"
    set_title_style(slide.shapes.title, 30)

    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12.0), Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True

    m0 = yolo_base["metrics"]
    m1 = yolo_retrain["metrics"]
    tm = temporal["best_metrics"]

    lines = [
        "YOLO-nano 재학습 결과",
        (
            f"- e3/f0.02: precision {m0['metrics/precision(B)']:.4f}, recall {m0['metrics/recall(B)']:.4f}, "
            f"mAP50-95 {m0['metrics/mAP50-95(B)']:.4f}"
        ),
        (
            f"- e20/f0.1: precision {m1['metrics/precision(B)']:.4f}, recall {m1['metrics/recall(B)']:.4f}, "
            f"mAP50-95 {m1['metrics/mAP50-95(B)']:.4f}"
        ),
        "",
        "Temporal GRU 결과",
        (
            f"- val_context_acc {tm['val_context_acc']:.4f}, val_boundary_f1 {tm['val_boundary_f1']:.4f}, "
            f"val_uncertainty_mae {tm['val_uncertainty_mae']:.4f}"
        ),
        "- 해석: context는 안정적, boundary precision 개선 필요",
    ]

    for idx, line in enumerate(lines):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = line
        para.level = 0

    set_body_style(tf, 18)


def add_race_result_slide(prs: Presentation, summary: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "race.mp4 실시간 실행 결과"
    set_title_style(slide.shapes.title, 30)

    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()

    bullets = [
        f"run_id: {summary['run_id']}",
        f"처리 프레임: {summary['frames']}",
        f"평균 FPS: {summary['avg_fps']}",
        f"총 검출 수: {summary.get('detection_count_total', 0)}",
        f"프레임당 평균 검출 수: {summary.get('detection_count_avg_per_frame', 0)}",
        f"검출 발생 프레임 비율: {summary.get('detection_nonzero_frame_ratio', 0)}",
        (
            "저장 정책 분포: "
            f"keyframe={summary['policy_counts'].get('keyframe', 0)}, "
            f"clip={summary['policy_counts'].get('clip', 0)}, "
            f"discard={summary['policy_counts'].get('discard', 0)}"
        ),
        f"이벤트 수: {summary['events']} (현재 threshold 기준)",
    ]

    for idx, item in enumerate(bullets):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = item
        para.level = 0

    set_body_style(tf, 18)


def add_screenshot_slide(prs: Presentation, screenshot_path: Path) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "실행 화면 예시"
    set_title_style(slide.shapes.title, 30)

    if screenshot_path.exists():
        slide.shapes.add_picture(str(screenshot_path), Inches(0.6), Inches(1.3), width=Inches(8.4))

    note = slide.shapes.add_textbox(Inches(9.2), Inches(1.6), Inches(3.8), Inches(4.2))
    tf = note.text_frame
    tf.word_wrap = True
    msgs = [
        "화면에서 확인되는 항목",
        "- YOLO 박스/클래스/confidence",
        "- context top-3 확률",
        "- score/context/event/FPS",
        "",
        "현재 관찰",
        "- 인식은 정상",
        "- 이벤트 트리거는 보수적",
    ]
    for idx, m in enumerate(msgs):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = m
    set_body_style(tf, 16)


def add_conclusion_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "결론 및 다음 단계"
    set_title_style(slide.shapes.title, 30)

    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    bullets = [
        "현재 상태: 통합 파이프라인/실시간 처리/결과 저장 모두 정상 동작",
        "강점: 구조 일관성, 데이터 축적 경로, 실시간성 확보",
        "부족한 점: 이벤트 경계 트리거 민감도, boundary precision",
        "다음 액션 1: race 전용 threshold 튜닝(clip 이벤트 활성화)",
        "다음 액션 2: boundary head 학습 샘플링/가중치 개선",
        "다음 액션 3: 정확도 평가(mAP)와 데모 평가를 분리 운영",
    ]
    for idx, item in enumerate(bullets):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.text = item
    set_body_style(tf, 18)


def build_talk_track(path: Path, summary: dict) -> None:
    lines = [
        "# 발표 스크립트",
        "",
        "## 1. 프로젝트 목적",
        "이 발표는 race.mp4를 이용해 비전인식-맥락 연계형 파이프라인이 실제로 동작하는지 검증한 결과입니다.",
        "",
        "## 2. 핵심 구조",
        "YOLO로 프레임 특징을 만들고, GRU로 시간 맥락을 해석한 뒤, scoring/curation/export까지 한 경로로 연결했습니다.",
        "",
        "## 3. 학습 결과",
        "YOLO 재학습으로 detector 성능을 개선했고, temporal은 context 분류가 안정적이지만 boundary precision은 추가 개선이 필요합니다.",
        "",
        "## 4. race.mp4 실시간 결과",
        (
            f"이번 실행에서 {summary['frames']}프레임을 처리했고 평균 FPS는 {summary['avg_fps']}입니다. "
            f"프레임당 평균 검출 수는 {summary.get('detection_count_avg_per_frame', 0)}입니다."
        ),
        "파이프라인은 정상 동작하지만 이벤트 트리거는 현재 threshold가 보수적이라 0건입니다.",
        "",
        "## 5. 결론",
        "현재 단계는 구조 검증과 실시간 운용성 확인에 성공했습니다. 다음 단계는 이벤트 민감도와 boundary 정밀도 개선입니다.",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out_dir = root / "artifacts" / "presentations"
    out_dir.mkdir(parents=True, exist_ok=True)

    summary, summary_path = latest_race_summary(root)
    yolo_base, yolo_retrain = load_yolo_metrics(root)
    temporal = load_temporal_metrics(root)

    screenshot = root / str(summary.get("screenshot", ""))
    date_tag = datetime.now().strftime("%Y%m%d")
    ppt_path = out_dir / f"vcp_race_demo_report_{date_tag}.pptx"
    talk_path = out_dir / "vcp_race_demo_talk_track.md"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs, datetime.now().strftime("%Y-%m-%d"))
    add_bullet_slide(
        prs,
        "문제 정의",
        [
            "목표: 실시간 카메라 입력에서 인식-맥락-정제-저장 파이프라인 검증",
            "환경: 학습(RTX 3080 Ti) / 배포 타깃(Jetson Orin Nano 8GB) 분리",
            "검증 데이터: race.mp4 (실시간 시각화 + 산출물 생성)",
        ],
    )
    add_bullet_slide(
        prs,
        "시스템 구조",
        [
            "Spatial: YOLO-nano 기반 특징 벡터화",
            "Temporal: GRU 기반 context/boundary/uncertainty 추정",
            "Scoring: context/boundary/uncertainty/novelty 결합",
            "Curation: clip/keyframe/discard 정책",
            "Export: frame/event/aligned dataset 저장",
        ],
    )
    add_metrics_slide(prs, yolo_base, yolo_retrain, temporal)
    add_race_result_slide(prs, summary)
    add_screenshot_slide(prs, screenshot)
    add_conclusion_slide(prs)

    prs.save(ppt_path)
    build_talk_track(talk_path, summary)

    manifest = {
        "ppt_path": str(ppt_path),
        "talk_track": str(talk_path),
        "summary_source": str(summary_path),
    }
    (out_dir / "presentation_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
