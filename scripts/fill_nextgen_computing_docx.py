from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = sorted(ROOT.glob("2026_*_v1.docx"))[0]
BACKUP_PATH = ROOT / "paper" / "manuscript" / f"{DOCX_PATH.stem}_original_backup.docx"
FILLED_COPY_PATH = ROOT / "paper" / "manuscript" / f"{DOCX_PATH.stem}_filled.docx"
ARCH_PATH = ROOT / "project_architecture.png"
SCREENSHOT_PATH = ROOT / "artifacts" / "screenshots" / "race_demo_20260319_133858.jpg"


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def clear_paragraph(paragraph) -> None:
    element = paragraph._element
    for child in list(element):
        element.remove(child)


def set_text(paragraph, text: str, *, bold: bool | None = None) -> None:
    sample = paragraph.runs[0] if paragraph.runs else None
    font_name = sample.font.name if sample and sample.font.name else None
    font_size = sample.font.size if sample and sample.font.size else None
    font_bold = sample.font.bold if sample and sample.font.bold is not None else None
    alignment = paragraph.alignment
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if font_name:
        run.font.name = font_name
    if font_size:
        run.font.size = font_size
    if bold is None:
        if font_bold is not None:
            run.font.bold = font_bold
    else:
        run.font.bold = bold
    paragraph.alignment = alignment


def set_cell_text(cell, text: str, *, bold: bool | None = None) -> None:
    sample = cell.paragraphs[0]
    font_name = sample.runs[0].font.name if sample.runs and sample.runs[0].font.name else None
    font_size = sample.runs[0].font.size if sample.runs and sample.runs[0].font.size else None
    font_bold = (
        sample.runs[0].font.bold
        if sample.runs and sample.runs[0].font.bold is not None
        else None
    )
    alignment = sample.alignment

    cell.text = text
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        run = paragraph.runs[0]
        if font_name:
            run.font.name = font_name
        if font_size:
            run.font.size = font_size
        if bold is None:
            if font_bold is not None:
                run.font.bold = font_bold
        else:
            run.font.bold = bold
    paragraph.alignment = alignment


def style_like(paragraph, source_paragraph) -> None:
    paragraph.style = source_paragraph.style
    paragraph.alignment = source_paragraph.alignment
    if paragraph.runs and source_paragraph.runs:
        src = source_paragraph.runs[0]
        dst = paragraph.runs[0]
        if src.font.name:
            dst.font.name = src.font.name
        if src.font.size:
            dst.font.size = src.font.size
        if src.font.bold is not None:
            dst.font.bold = src.font.bold


def insert_table_after(doc: Document, paragraph, rows: int, cols: int, style_name: str | None = None):
    table = doc.add_table(rows=rows, cols=cols)
    if style_name:
        table.style = style_name
    paragraph._p.addnext(table._tbl)
    return table


def fill_table(table, headers: list[str], rows: list[list[str]]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)

    while len(table.columns) < len(headers):
        table.add_column(Inches(1.0))

    header_row = table.rows[0]
    for idx, header in enumerate(headers):
        set_text(header_row.cells[idx].paragraphs[0], header, bold=True)

    if not rows:
        return

    if len(table.rows) < 2:
        table.add_row()

    first = rows[0]
    while len(header_row.cells) < len(first):
        table.add_column(Inches(1.0))

    first_row = table.rows[1]
    for idx, value in enumerate(first):
        set_text(first_row.cells[idx].paragraphs[0], value)

    for row_values in rows[1:]:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            set_text(row.cells[idx].paragraphs[0], value)


def prepare_table(table, col_count: int) -> None:
    while len(table.columns) < col_count:
        table.add_column(Inches(1.0))
    while len(table.rows) < 2:
        table.add_row()


def fill_title_table(doc: Document) -> None:
    table = doc.tables[0]
    set_cell_text(table.rows[0].cells[2], "멀티모달 데이터 선별을 위한 CNN-시계열 융합 모델 기법")
    set_cell_text(table.rows[2].cells[2], "CNN-Temporal Fusion Model for Multimodal Data Selection")
    set_cell_text(table.rows[4].cells[2], "작성자 정보 추후 입력")
    set_cell_text(table.rows[6].cells[2], "소속 및 이메일 정보 추후 입력")
    set_cell_text(
        table.rows[8].cells[1],
        (
            "본 논문은 실시간 카메라 입력으로부터 학습 가치가 높은 장면을 선별하기 위한 CNN-시계열 융합 모델을 제안한다. "
            "제안 구조는 YOLO 계열 공간 특징 추출, 검출 통계의 고정 길이 벡터화, GRU 기반 맥락 인코딩, 이벤트 점수화, 중요도 기반 저장으로 구성된다. "
            "실제 차량 영상 기반 구현 결과와 합성 10맥락 시계열 실험을 통해 구조적 타당성을 검증하였다."
        ),
    )
    set_cell_text(
        table.rows[9].cells[1],
        "키워드: 멀티모달 데이터 선별, YOLO, GRU, 시계열 맥락 인식, 엣지 AI",
    )


def fill_body(doc: Document, result_table) -> None:
    paragraphs = doc.paragraphs

    intro_heading = paragraphs[1]
    intro_body = paragraphs[2]
    related_heading = paragraphs[4]
    related_body = paragraphs[5]
    figure_slot = paragraphs[6]
    figure_caption = paragraphs[7]
    method_heading = paragraphs[9]
    dataset_heading = paragraphs[10]
    dataset_body = paragraphs[11]
    env_heading = paragraphs[13]
    env_body = paragraphs[14]
    result_body = paragraphs[15]
    table_caption = paragraphs[16]
    conclusion_heading = paragraphs[18]
    conclusion_body = paragraphs[19]
    ack_heading = paragraphs[21]
    corr_author = paragraphs[22]
    ack_body = paragraphs[23]
    ref_style_source = paragraphs[26]

    remove_list = list(doc.paragraphs[25:])
    for paragraph in remove_list:
        remove_paragraph(paragraph)

    result_table_style = result_table.style

    set_text(intro_heading, "1. 서론")
    set_text(
        intro_body,
        (
            "실환경 영상에서 학습 가치가 높은 구간만 선별하는 기술은 멀티모달 데이터셋 구축 비용을 줄이는 핵심 요소이다. "
            "그러나 단일 프레임 객체 검출만으로는 급정거, 돌발 보행자, 이륜차 끼어듦과 같은 시간적 맥락을 충분히 설명하기 어렵다. "
            "이에 본 연구는 경량 CNN 기반 공간 특징과 GRU 기반 시계열 맥락 인코딩을 결합하여, 실시간 비전 입력을 이벤트 중심 학습 샘플로 변환하는 구조를 구현하였다."
        ),
    )

    set_text(related_heading, "1.1 관련연구")
    set_text(
        related_body,
        (
            "YOLO 계열과 MobileNet 계열은 엣지 환경에서 널리 활용되는 경량 시각 인식 구조이며 [1,2], GRU는 비교적 낮은 연산량으로 연속 관측의 맥락을 모델링할 수 있는 대표 시계열 인코더이다 [3]. "
            "또한 TSM은 효율적 시간 모델링의 대안으로 알려져 있고 [4], RT-2와 OpenVLA는 정렬된 시각-행동 데이터의 중요성을 보여주었다 [5,6]. "
            "본 연구는 이러한 흐름 위에서 객체 검출, 시계열 맥락 추론, 저장 정책을 하나의 파이프라인으로 통합하는 구현에 초점을 둔다."
        ),
    )

    clear_paragraph(figure_slot)
    figure_slot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_slot.add_run().add_picture(str(ARCH_PATH), width=Inches(4.7))
    set_text(figure_caption, "그림 1. CNN-시계열 융합 기반 멀티모달 데이터 선별 파이프라인")

    set_text(method_heading, "2. 연구 방법론")
    set_text(dataset_heading, "2.1 데이터 및 문제 정의")
    set_text(
        dataset_body,
        (
            "실험 데이터는 차량 외부 주행 공개 데이터셋과 검증용 영상 `race.mp4`, `stopcar.mp4`로 구성하였다. 공간 인코더는 객체 라벨이 포함된 영상 데이터를 사용해 학습하였고, 시계열 인코더는 프레임별 검출 결과를 고정 길이 시퀀스로 변환한 processed dataset과 합성 10맥락 벡터 데이터셋으로 학습하였다."
        ),
    )

    set_text(env_heading, "2.2 모델 구조")
    set_text(
        env_body,
        (
            "공간 인코더는 YOLO 기반 객체 검출기의 출력으로부터 검출 수, confidence 통계, 박스 면적 비율, 클래스 비율, ROI 위험도, motion/looming 지표를 추출해 고정 길이 벡터로 구성한다. "
            "시계열 인코더는 최근 T개 프레임 벡터를 입력받는 GRU 구조이며, context, boundary, uncertainty head를 분리하여 맥락과 이벤트 전환을 동시에 예측한다."
        ),
    )

    param_heading = result_body.insert_paragraph_before("2.3 학습 및 실행 파라미터")
    style_like(param_heading, dataset_heading)
    param_body = result_body.insert_paragraph_before(
        "표 1은 본 연구에서 사용한 핵심 파라미터를 정리한 것이다. 학습 경로와 배포 경로는 분리하며, 본 논문에서는 데스크탑 학습 기준 결과를 사용하였다."
    )
    style_like(param_body, dataset_body)
    param_caption = result_body.insert_paragraph_before("표 1. 핵심 모델 및 학습 파라미터")
    style_like(param_caption, figure_caption)
    param_table = insert_table_after(doc, param_caption, rows=2, cols=3, style_name=result_table_style)
    prepare_table(param_table, 3)
    fill_table(
        param_table,
        headers=["구분", "설정값", "설명"],
        rows=[
            ["입력 벡터", "16차원", "검출 통계 + 위험 보조 특징"],
            ["시퀀스 길이", "8 frame", "GRU 입력 윈도"],
            ["GRU hidden", "96", "합성 10맥락 학습 설정"],
            ["Batch / Epoch", "512 / 12", "합성셋 GPU 학습 설정"],
            ["학습 장치", "RTX 3080 Ti / CUDA 11.8", "데스크탑 실험 환경"],
            ["저장 정책", "high=clip, mid=keyframe", "중요도 기반 자동 정제"],
        ],
    )

    result_heading = result_body.insert_paragraph_before("3. 실험 환경 및 결과")
    style_like(result_heading, intro_heading)
    result_sub_heading = result_body.insert_paragraph_before("3.1 실제 데이터 기반 결과")
    style_like(result_sub_heading, dataset_heading)
    set_text(
        result_body,
        (
            "실제 차량 영상 기반 실험에서 YOLO 공간 인코더와 GRU 시계열 인코더의 기본 동작을 검증하였다. YOLO 재학습 결과는 초기 기준 모델 대비 전반적으로 향상되었고, 실제 processed sequence에 대한 GRU 학습에서는 높은 context accuracy를 확인하였다. 다만 boundary F1은 상대적으로 낮아 이벤트 경계 표본 확장이 필요하다."
        ),
    )

    set_text(table_caption, "표 2. 실제 데이터 및 런타임 실험 결과")
    prepare_table(result_table, 3)
    fill_table(
        result_table,
        headers=["실험", "설정", "결과"],
        rows=[
            ["YOLO 공간 인코더", "차량 외부 영상 재학습", "precision 0.5804 / recall 0.4579 / mAP50 0.4763 / mAP50-95 0.2904"],
            ["실제 sequence GRU", "processed vehicle dataset", "val_context_acc 0.9662 / boundary_F1 0.2010 / uncertainty_MAE 0.0567"],
            ["실시간 데모", "race.mp4, 300 frame", "avg_fps 37.09, event 0건(기존 threshold 기준)"],
            ["급정지 정성 검증", "stopcar.mp4", "threshold 0.51 / min_gap 150에서 핵심 이벤트 1건 분리"],
        ],
    )

    synth_heading = conclusion_heading.insert_paragraph_before("3.2 합성 10맥락 학습 결과")
    style_like(synth_heading, dataset_heading)
    synth_body = conclusion_heading.insert_paragraph_before(
        "실제 데이터의 희소한 이벤트 분포를 보완하기 위해 10개 맥락, 6,400개 시계열 벡터로 구성된 합성 데이터셋을 생성하였다. 합성 규칙은 전방 차량 접근, 급제동 위험, 보행자 측면 대기/횡단/돌발 진입, 정지 이륜차, 이륜차 끼어듦 등을 포함한다."
    )
    style_like(synth_body, dataset_body)
    synth_caption = conclusion_heading.insert_paragraph_before("표 3. 합성 10맥락 GRU 학습 결과")
    style_like(synth_caption, figure_caption)
    synth_table = insert_table_after(doc, synth_caption, rows=2, cols=3, style_name=result_table_style)
    prepare_table(synth_table, 3)
    fill_table(
        synth_table,
        headers=["항목", "설정", "결과"],
        rows=[
            ["합성 데이터셋", "10 context / 6,400 sample", "train 5,120 / val 1,280"],
            ["GRU 설정", "window 8 / hidden 96 / batch 512 / epoch 12", "device cuda:0"],
            ["검증 성능", "규칙 기반 합성셋", "val_context_acc 1.0000 / boundary_F1 1.0000 / uncertainty_MAE 0.038586"],
            ["해석", "구조 검증 중심", "실제 일반화 성능이 아니라 분리 가능성 확인 결과"],
        ],
    )

    visual_heading = conclusion_heading.insert_paragraph_before("3.3 프로그램 실행 화면 및 정성 분석")
    style_like(visual_heading, dataset_heading)
    visual_body = conclusion_heading.insert_paragraph_before(
        "그림 2는 실시간 데모 실행 화면 예시이다. 화면에는 객체 검출 결과와 함께 현재 파이프라인이 출력한 점수 및 이벤트 관련 시각 정보가 반영되며, 이는 추론 결과의 해석 가능성과 검수 기반 후속 재학습 루프에 활용될 수 있다."
    )
    style_like(visual_body, dataset_body)
    if SCREENSHOT_PATH.exists():
        screenshot_para = conclusion_heading.insert_paragraph_before("")
        screenshot_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        screenshot_para.add_run().add_picture(str(SCREENSHOT_PATH), width=Inches(4.8))
        screenshot_caption = conclusion_heading.insert_paragraph_before("그림 2. 프로그램 실행 화면 예시")
        style_like(screenshot_caption, figure_caption)

    set_text(conclusion_heading, "4. 결론")
    set_text(
        conclusion_body,
        (
            "본 논문은 CNN 기반 공간 특징 추출과 GRU 기반 시계열 맥락 인코딩을 결합한 멀티모달 데이터 선별 구조를 구현하고, 실제 차량 영상과 합성 10맥락 데이터셋을 통해 동작 가능성을 검증하였다. 실제 데이터 실험에서는 실시간 처리와 기본 맥락 인코딩이 가능함을 확인하였고, 합성 데이터 실험에서는 다중 맥락 분류 구조가 안정적으로 학습될 수 있음을 보였다. 향후에는 Jetson Orin Nano 실기기 검증, 실제 IMU/차량 상태 센서 정렬, human-in-the-loop 재학습을 통해 실제 데이터 일반화 성능을 확장 검증할 예정이다."
        ),
    )

    set_text(ack_heading, "Acknowledgement")
    set_text(corr_author, "* 교신저자 정보는 최종 제출본에서 반영 예정")
    set_text(ack_body, "본 문서는 현재 구현 완료 범위와 예비 실험 결과를 기준으로 작성하였다.")

    ref_heading = doc.add_paragraph("참 고 문 헌")
    style_like(ref_heading, intro_heading)
    references = [
        '[1] J. Redmon et al., "You Only Look Once: Unified, Real-Time Object Detection," CVPR, 2016.',
        '[2] A. Howard et al., "Searching for MobileNetV3," arXiv preprint arXiv:1905.02244, 2019.',
        '[3] K. Cho et al., "Learning Phrase Representations using RNN Encoder-Decoder for Statistical Machine Translation," EMNLP, 2014.',
        '[4] J. Lin et al., "TSM: Temporal Shift Module for Efficient Video Understanding," ICCV, 2019.',
        '[5] A. Brohan et al., "RT-2: Vision-Language-Action Models Transfer Web Knowledge to Robotic Control," CoRL, 2023.',
        '[6] M. Kim et al., "OpenVLA: An Open-Source Vision-Language-Action Model," arXiv preprint arXiv:2406.09246, 2024.',
    ]
    for reference in references:
        new_p = doc.add_paragraph(reference)
        style_like(new_p, ref_style_source)


def main() -> None:
    BACKUP_PATH.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    source_path = BACKUP_PATH if BACKUP_PATH.exists() else DOCX_PATH
    doc = Document(source_path)
    result_table = doc.tables[1]
    fill_title_table(doc)
    fill_body(doc, result_table=result_table)
    doc.save(DOCX_PATH)
    shutil.copy2(DOCX_PATH, FILLED_COPY_PATH)

    print(f"updated: {DOCX_PATH}")
    print(f"backup: {BACKUP_PATH}")
    print(f"copy: {FILLED_COPY_PATH}")


if __name__ == "__main__":
    main()
